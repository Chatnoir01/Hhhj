import csv
import re
from pathlib import Path
from openpyxl import load_workbook

USERNAME_RE = re.compile(r"^[A-Za-z0-9_]{5,32}$")


def normalize_username(value: object) -> str | None:
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    raw = raw.replace("https://t.me/", "").replace("http://t.me/", "").replace("t.me/", "")
    raw = raw.split("?", 1)[0].split("/", 1)[0].lstrip("@").strip()
    return raw.lower() if USERNAME_RE.fullmatch(raw) else None


def _dedupe(values):
    seen = set()
    out = []
    for value in values:
        username = normalize_username(value)
        if username and username not in seen:
            seen.add(username)
            out.append(username)
    return out


def load_usernames(path: str | Path) -> list[str]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return _dedupe(path.read_text(encoding="utf-8").splitlines())
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = csv.reader(handle)
            return _dedupe(cell for row in rows for cell in row)
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        return _dedupe(cell.value for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    raise ValueError("Formats acceptés: .txt, .csv, .xlsx")
